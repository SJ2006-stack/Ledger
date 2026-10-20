"""Seed, save, and independently read Finance_AP_Ledger.xlsx.

Saving here is the spreadsheet application's persistence. The computer-use
agent reaches it by clicking Save in the workbook page. Callers that verify
a task only read.
"""

from __future__ import annotations

import threading
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from app.config import ROOT
from app.cua.schemas import COLUMNS, LEDGER_STATUSES, SHEET_NAMES

DEFAULT_WORKBOOK = ROOT / "backend" / "data" / "Finance_AP_Ledger.xlsx"
_LOCK = threading.Lock()

MONEY_COLUMNS = {
    "Invoice_Ledger": {"Subtotal", "Tax Amount", "Total"},
    "Purchase_Orders": {"Unit Price"},
}
HEADER_FILL = PatternFill("solid", fgColor="16382C")
HEADER_FONT = Font(color="F4F7F5", bold=True)


def seed_rows() -> dict[str, list[list[str]]]:
    """Mirror the SQLite seed. INV-2001 is intentionally absent. INV-3003 is the duplicate fixture."""
    return {
        "Vendors": [
            ["1", "Northwind Supplies", "27AABCU9603R1ZM", "Yes", "Net 30"],
            ["2", "Harbor Components", "29AABCH1234F1Z5", "Yes", "Net 30"],
            ["3", "Apex Industrial", "07AADCA1111B1Z8", "Yes", "Net 30"],
            ["4", "Helios Metals", "24AABCH5555A1Z1", "Yes", "Net 30"],
            ["5", "Helios Metal Works", "24AABCH6666A1Z2", "Yes", "Net 30"],
        ],
        "Purchase_Orders": [
            ["PO-1001", "Northwind Supplies", "SKU-WA", "Widget A", "100", "50.00", "INR"],
            ["PO-2002", "Harbor Components", "SKU-BB", "Bolt B", "100", "10.00", "INR"],
            ["PO-3003", "Apex Industrial", "SKU-CX", "Cable X", "10", "200.00", "INR"],
            ["PO-5005", "Helios Metals", "SKU-HM", "Plate HM", "5", "1000.00", "INR"],
            ["PO-5006", "Helios Metal Works", "SKU-HW", "Plate HW", "5", "1100.00", "INR"],
        ],
        "Goods_Receipts": [
            ["GRN-1001", "PO-1001", "SKU-WA", "100", "2026-09-01"],
            ["GRN-2002", "PO-2002", "SKU-BB", "90", "2026-09-01"],
            ["GRN-3003", "PO-3003", "SKU-CX", "10", "2026-09-01"],
            ["GRN-5005", "PO-5005", "SKU-HM", "5", "2026-09-01"],
            ["GRN-5006", "PO-5006", "SKU-HW", "5", "2026-09-01"],
        ],
        "Invoice_Ledger": [
            [
                "INV-3003",
                "Apex Industrial",
                "07AADCA1111B1Z8",
                "2026-09-15",
                "PO-3003",
                "GRN-3003",
                "2000.00",
                "18",
                "360.00",
                "2360.00",
                "NEEDS_APPROVAL",
                "",
                "apex-inv-3003.pdf",
                "seed",
            ]
        ],
    }


def ensure_workbook(path: Path | None = None) -> Path:
    target = path or DEFAULT_WORKBOOK
    with _LOCK:
        if not target.exists():
            write_sheets(target, seed_rows())
    return target


def reset_workbook(path: Path) -> None:
    with _LOCK:
        write_sheets(path, seed_rows())


def write_sheets(path: Path, sheets: dict[str, list[list[str]]]) -> None:
    cleaned = _validate(sheets)
    path.parent.mkdir(parents=True, exist_ok=True)
    book = Workbook()
    first = True
    for name in SHEET_NAMES:
        grid = book.active if first else book.create_sheet(name)
        first = False
        grid.title = name
        columns = COLUMNS[name]
        grid.append(list(columns))
        for cell in grid[1]:
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
        for row in cleaned[name]:
            grid.append(row)
        for index, column in enumerate(columns, start=1):
            letter = get_column_letter(index)
            grid.column_dimensions[letter].width = max(14, len(column) + 2)
        last_row = max(2, 1 + len(cleaned[name]))
        last_col = get_column_letter(len(columns))
        table = Table(displayName=name, ref=f"A1:{last_col}{max(last_row, 2)}")
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        grid.add_table(table)
        _format_numbers(grid, name)
    temporary = path.with_suffix(".tmp.xlsx")
    book.save(temporary)
    temporary.replace(path)


def read_sheets(path: Path) -> dict[str, list[dict[str, str]]]:
    if not path.exists():
        raise FileNotFoundError(path)
    book = load_workbook(path, data_only=True, read_only=True)
    try:
        loaded: dict[str, list[dict[str, str]]] = {}
        for name in SHEET_NAMES:
            grid = book[name]
            rows = list(grid.iter_rows(values_only=True))
            header = [str(cell or "") for cell in rows[0]] if rows else []
            parsed: list[dict[str, str]] = []
            money = MONEY_COLUMNS.get(name, set())
            for raw in rows[1:]:
                values = [_cell(value, header[index] in money if index < len(header) else False) for index, value in enumerate(raw)]
                if not any(values):
                    continue
                parsed.append({column: values[index] if index < len(values) else "" for index, column in enumerate(header)})
            loaded[name] = parsed
        return loaded
    finally:
        book.close()


def sheets_as_rows(path: Path) -> dict[str, list[list[str]]]:
    loaded = read_sheets(path)
    return {
        name: [[row.get(column, "") for column in COLUMNS[name]] for row in loaded[name]]
        for name in SHEET_NAMES
    }


def ledger_rows(path: Path, invoice_id: str) -> list[dict[str, str]]:
    return [row for row in read_sheets(path)["Invoice_Ledger"] if row.get("Invoice ID") == invoice_id]


def verify_ledger(
    path: Path,
    *,
    invoice_id: str,
    expect_row: bool,
    intended: dict[str, str] | None = None,
    source_pdf: str = "",
) -> dict:
    """Read the xlsx from disk. This does not ask the computer-use controller what it thinks happened."""
    return verify_rows(
        read_sheets(path)["Invoice_Ledger"],
        invoice_id=invoice_id,
        expect_row=expect_row,
        intended=intended,
        source_pdf=source_pdf,
    )


def verify_rows(
    rows: list[dict[str, str]],
    *,
    invoice_id: str,
    expect_row: bool,
    intended: dict[str, str] | None = None,
    source_pdf: str = "",
) -> dict:
    if invoice_id:
        found = [row for row in rows if row.get("Invoice ID") == invoice_id]
    elif source_pdf:
        found = [row for row in rows if row.get("Source PDF") == source_pdf]
    else:
        found = []
    mismatches: list[str] = []
    if expect_row:
        if len(found) != 1:
            mismatches.append(f"expected one row for {invoice_id or source_pdf}, found {len(found)}")
        elif intended:
            actual = found[0]
            for column, wanted in intended.items():
                if column == "Processed At":
                    if not actual.get(column):
                        mismatches.append("Processed At is empty")
                    continue
                if _same(actual.get(column, ""), wanted):
                    continue
                mismatches.append(f"{column}: workbook has {actual.get(column, '')!r}, intended {wanted!r}")
    elif invoice_id == "INV-3003":
        if len(found) != 1:
            mismatches.append(f"duplicate fixture should stay a single row, found {len(found)}")
        elif found[0].get("Vendor") != "Apex Industrial":
            mismatches.append("The seeded Apex row was changed.")
    elif len(found) != 0:
        mismatches.append(f"expected no ledger row, found {len(found)}")
    return {
        "ok": not mismatches,
        "mismatches": mismatches,
        "rows": found,
        "checked_at": datetime.now(timezone.utc).isoformat(),
    }


def _validate(sheets: dict[str, list[list[str]]]) -> dict[str, list[list[str]]]:
    cleaned: dict[str, list[list[str]]] = {}
    for name in SHEET_NAMES:
        columns = COLUMNS[name]
        rows = sheets.get(name) or []
        kept: list[list[str]] = []
        seen: set[str] = set()
        for raw in rows:
            values = [_cell(raw[index] if index < len(raw) else "") for index in range(len(columns))]
            if not any(values):
                continue
            if name == "Invoice_Ledger":
                invoice_id = values[0]
                status = values[10]
                if invoice_id:
                    if invoice_id in seen:
                        raise ValueError(f"Duplicate invoice id {invoice_id}.")
                    seen.add(invoice_id)
                if status and status not in LEDGER_STATUSES:
                    raise ValueError(f"Status {status} is not allowed.")
            kept.append(values)
        cleaned[name] = kept
    return cleaned


def _format_numbers(grid, name: str) -> None:
    money_columns = {
        "Invoice_Ledger": {7, 9, 10},
        "Purchase_Orders": {6},
    }.get(name, set())
    for row in grid.iter_rows(min_row=2):
        for index, cell in enumerate(row, start=1):
            if index in money_columns and isinstance(cell.value, str):
                try:
                    cell.value = float(cell.value)
                    cell.number_format = '#,##0.00'
                except ValueError:
                    continue


def _cell(value: object, money: bool = False) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, (int, float)) and money:
        return f"{float(value):.2f}"
    if isinstance(value, float):
        return f"{value:.2f}"
    return str(value).strip()


def _same(actual: str, wanted: str) -> bool:
    if actual == wanted:
        return True
    try:
        return float(actual) == float(wanted)
    except ValueError:
        return False
